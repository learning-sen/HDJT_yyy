# -*- coding: utf-8 -*-
"""把 PyAether 教程 Markdown 转成 DOCX（供 LibreOffice 转 PDF）。"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1] if len(sys.argv) > 1 else 'PyAether新手速成教程.md'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'PyAether新手速成教程.docx'

BODY_FONT = 'Times New Roman'
CJK_BODY = '宋体'
CJK_HEAD = '微软雅黑'

doc = Document()

# 页面：A4，稍窄边距
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)

normal = doc.styles['Normal']
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), CJK_BODY)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.15


def set_cjk(run, font=None):
    run.font.name = font or BODY_FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), CJK_BODY)


def add_rich(par, text, bold=False, size=None, color=None, head=False):
    """处理 **粗体** 与 `代码`。"""
    parts = re.split(r'(\*\*.+?\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Consolas')
            r.font.size = Pt((size or 10.5) - 0.5)
            r.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
        else:
            r = par.add_run(part)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        if head:
            r.font.name = CJK_HEAD
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), CJK_HEAD)
        else:
            set_cjk(r)
    return par


def flush_table(rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.autofit = True
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        # 禁止行内跨页断裂
        trPr = t.rows[-1]._tr.get_or_add_trPr()
        cant = trPr.makeelement(qn('w:cantSplit'), {})
        trPr.append(cant)
        for j in range(ncol):
            txt = row[j] if j < len(row) else ''
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = (i == 0)
            add_rich(p, txt, bold=(i == 0), size=9)
    doc.add_paragraph()


def add_code(lines):
    """代码块：等宽字体 + 灰底 + 边框。"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), 'F5F5F5')
    tcPr.append(shd)
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(ln if ln else ' ')
        r.font.name = 'Consolas'
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Consolas')
        r.font.size = Pt(8.5)
    doc.add_paragraph()


lines = open(SRC, encoding='utf-8').read().splitlines()
tbl = []
i = 0
in_code = False
code_buf = []
while i < len(lines):
    line = lines[i].rstrip()
    stripped = line.strip()

    # 代码块
    if stripped.startswith('```'):
        if in_code:
            add_code(code_buf)
            code_buf = []
            in_code = False
        else:
            if tbl:
                flush_table(tbl)
                tbl = []
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    # 表格
    if stripped.startswith('|'):
        cells = [c.strip() for c in stripped.strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
            i += 1
            continue
        tbl.append(cells)
        i += 1
        continue
    if tbl:
        flush_table(tbl)
        tbl = []

    if not stripped:
        i += 1
        continue

    if stripped.startswith('---'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        i += 1
        continue

    m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', stripped)
    if m:
        alt, path = m.group(1), m.group(2)
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(path, width=Cm(13.5))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(8)
            r = cap.add_run(alt)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            set_cjk(r)
        except Exception as e:
            print('image failed:', path, e)
            p = doc.add_paragraph('【图片缺失：%s】' % path)
        i += 1
        continue

    m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
    if m:
        lvl = len(m.group(1))
        text = m.group(2)
        if lvl == 1:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(p, text, bold=True, size=20, head=True)
            p.paragraph_format.space_after = Pt(10)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10 if lvl == 2 else 6)
            p.paragraph_format.space_after = Pt(4)
            add_rich(p, text, bold=True, size=(14 if lvl == 2 else 11.5), head=True)
            if lvl == 2:
                p.runs[0].font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
        i += 1
        continue

    if stripped.startswith('> '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        add_rich(p, stripped[2:], size=10)
        for r in p.runs:
            r.italic = True
        i += 1
        continue

    m = re.match(r'^[-*]\s+(.*)$', stripped)
    if m:
        p = doc.add_paragraph(style='List Bullet')
        add_rich(p, m.group(1))
        p.paragraph_format.space_after = Pt(2)
        i += 1
        continue

    m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
    if m:
        # 手写编号，避免 Word 自动编号跨列表续号
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(2)
        add_rich(p, m.group(1) + '. ' + m.group(2))
        i += 1
        continue

    p = doc.add_paragraph()
    add_rich(p, stripped)
    i += 1

if in_code and code_buf:
    add_code(code_buf)
if tbl:
    flush_table(tbl)

doc.save(OUT)
print('saved', OUT)
