# -*- coding: utf-8 -*-
"""v1 / v2 报告自检：数字多重集、表格逐格、一级标题、代码块文本。"""
import re
import sys
from collections import Counter

from docx import Document

ROOT = "Submission/00_技术报告/"
V1 = ROOT + "技术报告_20261019.docx"
V2 = ROOT + "技术报告_20261019_v2.docx"


def all_text(doc):
    for p in doc.paragraphs:
        yield p.text
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield c.text


def num_tokens(doc):
    txt = "\n".join(all_text(doc))
    return Counter(re.findall(r"\d+(?:\.\d+)?", txt))


def h1s(doc):
    return [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]


def main():
    d1, d2 = Document(V1), Document(V2)
    ok = True

    n1, n2 = num_tokens(d1), num_tokens(d2)
    diff = (n1 - n2) + (n2 - n1)
    print("数字 token 总数 v1/v2:", sum(n1.values()), sum(n2.values()))
    print("数字差异:", dict(diff) if diff else "无")
    ok &= not diff

    print("表格数 v1/v2:", len(d1.tables), len(d2.tables))
    ok &= len(d1.tables) == len(d2.tables)
    cell_diff = 0
    for i, (a, b) in enumerate(zip(d1.tables, d2.tables)):
        ra, rb = len(a.rows), len(b.rows)
        ca = len(a.columns)
        if ra != rb or len(b.columns) != ca:
            print(f"  表 {i}: 行/列数不同 {ra}x{ca} vs {rb}x{len(b.columns)}")
            ok = False
            continue
        for r in range(ra):
            for c in range(ca):
                if a.rows[r].cells[c].text != b.rows[r].cells[c].text:
                    cell_diff += 1
                    if cell_diff <= 10:
                        print(f"  表 {i} [{r},{c}]: {a.rows[r].cells[c].text!r} -> {b.rows[r].cells[c].text!r}")
    print("单元格差异数:", cell_diff)
    ok &= cell_diff == 0

    h1a, h1b = h1s(d1), h1s(d2)
    print("一级标题数 v1/v2:", len(h1a), len(h1b))
    if h1a != h1b:
        ok = False
        for x, y in zip(h1a, h1b):
            if x != y:
                print("  H1 不同:", x, "->", y)
    else:
        print("一级标题文字: 完全一致")

    # 代码块（等宽 Consolas 段落）逐行比对
    def mono(doc):
        out = []
        for p in doc.paragraphs:
            for r in p.runs:
                if r.font.name == "Consolas":
                    out.append(p.text)
                    break
        return out

    m1, m2 = mono(d1), mono(d2)
    print("代码/等宽行数 v1/v2:", len(m1), len(m2), "一致" if m1 == m2 else "不一致")
    if m1 != m2:
        ok = False
        for a, b in zip(m1, m2):
            if a != b:
                print("  ", repr(a), "->", repr(b))
        print("  长度差:", len(m2) - len(m1))

    print("\n结论:", "全部通过" if ok else "存在差异")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
